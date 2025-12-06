import pandas as pd
from openpyxl import load_workbook
from pathlib import Path
import math

src_path = Path("мой тренер окт 15-30.xlsx")  // для загрузки в репозиторий я поменяла названия, это изначальные файлы
dst_path = Path("мой тренер окт 15-30 подсчитано.xlsx")
sheet_name = "итого"
meta_cols_count = 3
total_col_idx = 2

df = pd.read_excel(src_path, sheet_name=sheet_name)

trainer_cols = df.columns[meta_cols_count:]
total_col = df.columns[total_col_idx]

def _to_num(x):
    if pd.isna(x):
        return math.nan
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(",", ".")
    return pd.to_numeric(s, errors="coerce")

def compute_my_trainers(row):
    total = _to_num(row[total_col])
    if pd.isna(total) or total <= 0:
        return ""

    pairs = []
    for col in trainer_cols:
        num = _to_num(row[col])
        if pd.isna(num) or num <= 0:
            continue
        ratio = float(num) / float(total)
        if ratio >= 0.3:
            pairs.append((col, num))

    pairs.sort(key=lambda x: (-x[1], str(x[0])))
    return ", ".join([name for name, _ in pairs[:3]])

df["Мои тренеры"] = df.apply(compute_my_trainers, axis=1)

cols = list(df.columns)
col_to_move = "Мои тренеры"

insert_after = 2

cols.insert(insert_after + 1, cols.pop(cols.index(col_to_move)))

df = df[cols]

wb = load_workbook(src_path)
if sheet_name in wb.sheetnames:
    wb.remove(wb[sheet_name])
ws = wb.create_sheet(sheet_name)

for j, col in enumerate(df.columns, start=1):
    ws.cell(row=1, column=j, value=col)

for i, (_, row) in enumerate(df.iterrows(), start=2):
    for j, col in enumerate(df.columns, start=1):
        ws.cell(row=i, column=j, value=row[col])